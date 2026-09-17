from __future__ import annotations

import argparse
import re
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from classify_orders import PROJECT_ROOT, classify_row, latest_excel_file, load_rules, normalize_key


WORKSPACE_ROOT = PROJECT_ROOT.parent
DEFAULT_TEMPLATE_DIR = PROJECT_ROOT / "templates"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "data" / "output" / "발주서"

REMOVE_PRODUCT_WORDS = [
    "아메리칸플로어",
    "아메리칸 플로어",
    "사계절",
    "매트",
    "카페트",
    "방수",
    "워셔블",
    "거실",
    "침실",
    "먼지없는",
]

CAPRI_COLORS = [
    "마리나 블루",
    "제스트 옐로우",
    "몬테 그린",
    "글리오니 그레이",
]

RUNE_COLORS = [
    "미스트 그레이",
    "웜 베이지",
    "퓨어 아이보리",
]

CHANNEL_TAGS = [
    "오늘의집",
    "스마트스토어",
    "Cafe24",
    "Cafe24(신)",
    "유튜브쇼핑",
    "29CM",
]

LOGIS_EXCLUDE_KEYWORDS = [
    "코니",
    "레이지데이",
    "모던타임즈",
    "팜하우스",
    "팜마우스",
    "모니카",
    "카프리섬머",
    "카프리 썸머",
    "카프리 시어서커",
    "시어서커",
    "스노우리지",
    "자카드",
    "레트로가든",
    "허쉬",
    "정글드로잉",
    "웨이비리듬",
    "블로썸체크",
    "메리그리드",
    "베른",
    "페블",
    "멜로우데이",
]

RUNE_SIZE_SHAPE_HINTS = {
    "150x200": "사각형",
    "170x230": "사각형",
    "200x270": "사각형",
    "200x400": "사각형",
    "100x100": "원형",
    "150x150": "원형",
    "170x170": "원형",
    "200x200": "원형",
    "150x225": "타원형",
    "75x200": "러너",
    "45x65": "반원매트",
}

RUNE_ROUND_SIZES = ["100", "150", "170", "200"]
RUNE_OVAL_SIZES = {
    "100": "100x150",
    "150": "150x225",
    "200": "200x300",
}


def compact_text(value: Any) -> str:
    return re.sub(r"\s+", "", normalize_key(value)).lower()


def clean_delivery_message(value: Any) -> str:
    message = normalize_key(value)
    if not message:
        return ""

    def replace_bracket(match: re.Match[str]) -> str:
        content = match.group(1)
        compacted = compact_text(content)
        if any(compact_text(tag) in compacted for tag in CHANNEL_TAGS):
            return ""
        return match.group(0)

    message = re.sub(r"\[([^\]]+)\]", replace_bracket, message)
    for tag in CHANNEL_TAGS:
        message = message.replace(f"[{tag}]", "").replace(tag, "")
    return re.sub(r"\s+", " ", message).strip()


def extract_labeled_value(text: str, labels: list[str]) -> str:
    for label in labels:
        match = re.search(rf"{label}\s*[=:]\s*([^,/\n]+)", text, flags=re.IGNORECASE)
        if match:
            return clean_option_value(match.group(1))
        bracket_match = re.search(rf"\[{label}\]\s*([^\[]+)", text, flags=re.IGNORECASE)
        if bracket_match:
            return clean_option_value(bracket_match.group(1))
    return ""


def clean_option_value(value: Any) -> str:
    text = normalize_key(value)
    text = re.sub(r"^\d+\s*[.)]\s*", "", text)
    text = text.replace("선택", "").strip()
    return re.sub(r"\s+", " ", text).strip()


def clean_color_value(value: Any) -> str:
    text = clean_option_value(value)
    text = re.sub(r"\((?:사각|사각형|원형|타원|타원형|반원매트|반원|반원형|러너)\)", "", text)
    return re.sub(r"\s+", " ", text).strip()


def extract_known_color(text: str, colors: list[str]) -> str:
    compacted = compact_text(text)
    for color in colors:
        if compact_text(color) in compacted:
            return color
    return ""


def extract_product_color(row: pd.Series) -> str:
    product = normalize_key(row.get("상품명", ""))
    if "," not in product:
        return ""

    head = product.split(",")[0]
    head = re.sub(r"\[[^\]]+\]", " ", head)
    head = re.sub(r"\d{2,3}\s*[xX*×]\s*\d{2,3}\s*(?:cm)?", " ", head)

    for word in REMOVE_PRODUCT_WORDS:
        head = head.replace(word, " ")

    pattern = clean_product_pattern(row)
    for token in pattern.split():
        head = head.replace(token, " ")

    for shape in ["타원형", "타원", "반원형", "반원", "사각형", "사각", "원형", "러너"]:
        head = head.replace(shape, " ")

    return clean_color_value(head)


def extract_color(row: pd.Series) -> str:
    option = normalize_key(row.get("옵션", ""))
    product = normalize_key(row.get("상품명", ""))
    known_color = extract_known_color(f"{option} {product}", CAPRI_COLORS + RUNE_COLORS)
    color = extract_labeled_value(option, ["컬러", "색상"]) or extract_labeled_value(product, ["컬러", "색상"])
    return known_color or clean_color_value(color) or extract_product_color(row)


def is_rune_product(text: str) -> bool:
    compacted = compact_text(text)
    return "룬" in compacted and "울트라" in compacted


def has_size_number(source: str, size: str) -> bool:
    return bool(re.search(rf"(?<!\d){size}\s*(?:cm)?(?!\d)", source, flags=re.IGNORECASE))


def extract_rune_named_size(source: str) -> str:
    compacted_source = compact_text(source)
    if not is_rune_product(source):
        return ""

    if "반원매트" in compacted_source:
        return "45x65"
    if "러너" in compacted_source:
        return "75x200"

    if "타원형" in compacted_source or "타원" in compacted_source:
        for size, normalized_size in RUNE_OVAL_SIZES.items():
            if has_size_number(source, size):
                return normalized_size

    if "원형" in compacted_source:
        for size in RUNE_ROUND_SIZES:
            if has_size_number(source, size):
                return f"{size}x{size}"

    return ""


def extract_size(row: pd.Series) -> str:
    source = " ".join([normalize_key(row.get("옵션", "")), normalize_key(row.get("상품명", ""))])
    match = re.search(r"(\d{2,3})\s*(?:cm)?\s*[xX*×]\s*(\d{2,3})\s*(?:cm)?", source)
    if match:
        return f"{match.group(1)}x{match.group(2)}"

    return extract_rune_named_size(source)


def extract_shape(row: pd.Series) -> str:
    source = " ".join([normalize_key(row.get("옵션", "")), normalize_key(row.get("상품명", ""))])
    compacted_source = compact_text(source)
    labeled = extract_labeled_value(source, ["디자인", "타입", "형태", "모양"])
    shape_source = compact_text(labeled) or compacted_source

    shape_aliases = [
        ("타원형", ["타원형", "타원"]),
        ("반원매트", ["반원매트"]),
        ("반원형", ["반원형", "반원"]),
        ("사각형", ["사각형", "사각"]),
        ("원형", ["원형"]),
        ("러너", ["러너"]),
        ("도어매트", ["도어매트"]),
    ]
    for normalized_shape, aliases in shape_aliases:
        if any(alias in shape_source for alias in aliases):
            return normalized_shape

    if is_rune_product(source):
        inferred_shape = RUNE_SIZE_SHAPE_HINTS.get(extract_size(row))
        if inferred_shape:
            return inferred_shape

    product_compacted = compact_text(row.get("상품명", ""))
    if "버드가든" in product_compacted or "버드가든" in compacted_source:
        size = extract_size(row)
        if size:
            width, height = size.split("x")
            return "원형" if width == height else "사각형"

    return ""


def clean_product_pattern(row: pd.Series) -> str:
    raw_name = normalize_key(row.get("상품명", ""))
    compacted = compact_text(raw_name)

    new_product_patterns = [
        ("노마드발수", "노마드 발수 러그"),
        ("어반위브사이잘룩", "어반 위브 사이잘룩 러그"),
        ("빌라브리즈시어서커", "빌라 브리즈 시어서커 러그"),
        ("리틀포니랜드시어서커", "리틀 포니랜드 시어서커 러그"),
        ("블루밍데이지시어서커", "블루밍 데이지 시어서커 러그"),
        ("카프리섬머", "카프리 시어서커 러그"),
        ("카프리썸머", "카프리 시어서커 러그"),
        ("카프리시어서커", "카프리 시어서커 러그"),
        ("트윙클다이아", "트윙클 다이아 러그"),
    ]
    for keyword, product_name in new_product_patterns:
        if keyword in compacted:
            return product_name

    if "시어서커" in compacted:
        return "시어서커 러그"
    if "멜로우데이" in compacted:
        return "멜로우데이 러그"
    if "레이지데이" in compacted:
        return "레이지데이 러그"
    if "버드가든" in compacted:
        return "버드가든 러그"
    if "스윔" in compacted:
        return "스윔 러그"
    if "룬" in compacted and "울트라" in compacted:
        return "룬 울트라 퍼포먼스 러그"

    name = re.sub(r"\[[^\]]+\]", " ", raw_name)
    name = name.split(",")[0]
    name = re.sub(r"\d{2,3}\s*[xX*×]\s*\d{2,3}\s*(?:cm)?", " ", name)

    for word in REMOVE_PRODUCT_WORDS:
        name = name.replace(word, " ")

    name = re.sub(r"\s+", " ", name).strip(" -_/")
    name = re.sub(r"(?<!\s)러그", " 러그", name)
    name = re.sub(r"\s+", " ", name).strip()

    rug_index = name.find("러그")
    if rug_index >= 0:
        return name[: rug_index + len("러그")].strip()

    return name


def format_product_name(row: pd.Series) -> str:
    pattern = clean_product_pattern(row)
    color = extract_color(row)
    size = extract_size(row)
    shape = extract_shape(row)

    parts = [part for part in [pattern, color] if part]
    compacted_pattern = compact_text(pattern)
    shape_aware_patterns = [
        "레이지데이",
        "멜로우데이",
        "버드가든",
        "룬울트라퍼포먼스",
        "노마드발수",
        "어반위브사이잘룩",
    ]
    if shape and any(pattern_key in compacted_pattern for pattern_key in shape_aware_patterns):
        parts.append(shape)
    if size:
        parts.append(size)
    return " / ".join(parts)


def derive_mmdd(input_path: Path, orders: pd.DataFrame) -> str:
    match = re.search(r"(20\d{6})", input_path.name)
    if match:
        return datetime.strptime(match.group(1), "%Y%m%d").strftime("%m%d")

    if "주문일자" in orders.columns:
        dates = pd.to_datetime(orders["주문일자"], errors="coerce").dropna()
        if not dates.empty:
            return dates.max().strftime("%m%d")

    return datetime.now().strftime("%m%d")


def copy_cell_style(source: Cell, target: Cell) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy(source.alignment)
    if source.protection:
        target.protection = copy(source.protection)


def clear_data_rows(sheet: Worksheet, header_row: int = 1) -> list[Cell]:
    template_row_number = header_row + 1 if sheet.max_row > header_row else header_row
    template_cells = [copy(sheet.cell(template_row_number, col)) for col in range(1, sheet.max_column + 1)]
    if sheet.max_row > header_row:
        sheet.delete_rows(header_row + 1, sheet.max_row - header_row)
    return template_cells


def write_template_workbook(
    template_path: Path,
    output_path: Path,
    rows: list[dict[str, Any]],
    sheet_name: str | None = None,
    defaults: dict[str, Any] | None = None,
) -> None:
    workbook = load_workbook(template_path)
    sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    template_cells = clear_data_rows(sheet)
    headers = [normalize_key(sheet.cell(1, col).value) for col in range(1, sheet.max_column + 1)]
    defaults = defaults or {}

    for row_index, row_data in enumerate(rows, start=2):
        for col_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row_index, col_index)
            if col_index <= len(template_cells):
                copy_cell_style(template_cells[col_index - 1], cell)
            cell.value = row_data.get(header, defaults.get(header, ""))

    sheet.freeze_panes = "A2"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def row_text(row: pd.Series) -> str:
    return " ".join([normalize_key(row.get("상품명", "")), normalize_key(row.get("옵션", ""))])


def should_exclude_from_logis(row: pd.Series) -> tuple[bool, str]:
    text = row_text(row)
    compacted = compact_text(text)
    if "러그" in compacted:
        return True, "상품명/옵션에 러그 포함"
    for keyword in LOGIS_EXCLUDE_KEYWORDS:
        if compact_text(keyword) in compacted:
            return True, f"제외 브랜드: {keyword}"
    return False, ""


def build_dream_rows(orders: pd.DataFrame) -> list[dict[str, Any]]:
    rows = []
    for _, row in orders.iterrows():
        rows.append(
            {
                "수취인": normalize_key(row.get("수취인명", "")),
                "전화번호1": normalize_key(row.get("연락처1", "")),
                "전화번호2": normalize_key(row.get("연락처2", "")),
                "주소": normalize_key(row.get("주소", "")),
                "수량": row.get("내품수량", ""),
                "품명": format_product_name(row),
                "요청사항": clean_delivery_message(row.get("배송메시지", "")),
                "운송장번호": "",
                "주문번호": "",
                "송하인 전화번호": "070-4035-8217",
                "송하인": "American Floor",
            }
        )
    return rows


def build_hyekyung_rows(orders: pd.DataFrame) -> list[dict[str, Any]]:
    rows = []
    for _, row in orders.iterrows():
        rows.append(
            {
                "상품명": format_product_name(row),
                "수량": row.get("내품수량", ""),
                "수취인명": normalize_key(row.get("수취인명", "")),
                "수취인연락처1": normalize_key(row.get("연락처1", "")),
                "수취인연락처2": normalize_key(row.get("연락처2", "")),
                "배송지우편번호": normalize_key(row.get("우편번호", "")),
                "배송지": normalize_key(row.get("주소", "")),
                "구매자메모": clean_delivery_message(row.get("배송메시지", "")),
            }
        )
    return rows


def build_logis_rows(orders: pd.DataFrame, template_path: Path) -> list[dict[str, Any]]:
    template = load_workbook(template_path)
    sheet = template.worksheets[0]
    headers = [normalize_key(sheet.cell(1, col).value) for col in range(1, sheet.max_column + 1)]

    rows = []
    for _, row in orders.iterrows():
        item = {header: row.get(header, "") for header in headers}
        if "배송메시지" in item:
            item["배송메시지"] = clean_delivery_message(item["배송메시지"])
        rows.append(item)
    return rows


def export_supplier_orders(input_path: Path, template_dir: Path, output_dir: Path) -> dict[str, Path]:
    rules = load_rules()
    orders = pd.read_excel(input_path)
    orders["발주처"] = [classify_row(row, rules)[0] for _, row in orders.iterrows()]

    logis_excluded_rows = []
    logis_keep_mask = []
    for _, row in orders.iterrows():
        if row["발주처"] != "로지스존":
            logis_keep_mask.append(False)
            continue
        excluded, reason = should_exclude_from_logis(row)
        logis_keep_mask.append(not excluded)
        if excluded:
            excluded_row = row.copy()
            excluded_row["제외사유"] = reason
            logis_excluded_rows.append(excluded_row)

    mmdd = derive_mmdd(input_path, orders)
    supplier_dir = output_dir / mmdd

    dream_template = template_dir / "드림산업 발주서 양식.xlsx"
    logis_template = template_dir / "로지스존 발주서 양식.xlsx"
    hyekyung_template = template_dir / "혜경산업 발주서 양식.xlsx"

    outputs = {
        "드림산업": supplier_dir / f"드림산업_{mmdd}.xlsx",
        "로지스존": supplier_dir / f"로지스존_{mmdd}.xlsx",
        "혜경산업": supplier_dir / f"혜경산업_{mmdd}.xlsx",
        "메르시데코": supplier_dir / f"메르시데코_{mmdd}.xlsx",
    }

    write_template_workbook(
        dream_template,
        outputs["드림산업"],
        build_dream_rows(orders[orders["발주처"] == "드림산업"]),
        sheet_name="발주서",
    )
    write_template_workbook(
        hyekyung_template,
        outputs["혜경산업"],
        build_hyekyung_rows(orders[orders["발주처"] == "혜경산업"]),
    )
    write_template_workbook(
        hyekyung_template,
        outputs["메르시데코"],
        build_hyekyung_rows(orders[orders["발주처"] == "메르시데코"]),
    )
    write_template_workbook(
        logis_template,
        outputs["로지스존"],
        build_logis_rows(orders[logis_keep_mask], logis_template),
    )

    legacy_hyekyung_path = supplier_dir / f"세븐펫-{mmdd}-발주서.xlsx"
    if legacy_hyekyung_path.exists():
        legacy_hyekyung_path.unlink()

    if logis_excluded_rows:
        excluded_path = supplier_dir / f"검수필요_{mmdd}.xlsx"
        pd.DataFrame(logis_excluded_rows).to_excel(excluded_path, index=False)
        outputs["검수필요"] = excluded_path
    else:
        excluded_path = supplier_dir / f"검수필요_{mmdd}.xlsx"
        if excluded_path.exists():
            excluded_path.unlink()

    return outputs


def main() -> None:
    parser = argparse.ArgumentParser(description="사방넷 주문서를 발주처별 발주서 양식으로 나눕니다.")
    parser.add_argument("--input", type=Path, help="정리할 사방넷 원본 엑셀 파일")
    parser.add_argument("--input-dir", type=Path, default=PROJECT_ROOT / "data" / "input", help="최신 엑셀 파일을 찾을 입력 폴더")
    parser.add_argument("--template-dir", type=Path, default=DEFAULT_TEMPLATE_DIR, help="발주처별 양식 폴더")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR, help="발주서 출력 폴더")
    args = parser.parse_args()

    input_path = args.input if args.input else latest_excel_file(args.input_dir)
    outputs = export_supplier_orders(input_path, args.template_dir, args.output_dir)
    for supplier, path in outputs.items():
        print(f"{supplier}: {path}")


if __name__ == "__main__":
    main()
